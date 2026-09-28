"""Local xlsx viewer/editor for the Claude app browser pane.

Usage: python3 server.py <file.xlsx> [port]
Binds 127.0.0.1 only. Saves back through openpyxl, so data validation and
conditional formatting in the workbook are kept.
"""
import datetime
import html
import json
import re
import sys
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

import openpyxl
from openpyxl.utils import column_index_from_string, range_boundaries

PATH = sys.argv[1]
PORT = int(sys.argv[2]) if len(sys.argv) > 2 else 8765
DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")


def dropdowns(wb, ws):
    """Map (row, col) -> option list for list-type validations."""
    out = {}
    for dv in ws.data_validations.dataValidation:
        if dv.type != "list" or not dv.formula1:
            continue
        f = dv.formula1
        if f.startswith('"'):
            opts = f.strip('"').split(",")
        else:
            sheet, ref = f.split("!") if "!" in f else (ws.title, f)
            src = wb[sheet.strip("'")]
            opts = [c.value for row in src[ref.replace("$", "")] for c in row if c.value is not None]
        for rng in str(dv.sqref).split():
            c1, r1, c2, r2 = range_boundaries(rng)
            for r in range(r1, r2 + 1):
                for c in range(c1, c2 + 1):
                    out[(r, c)] = [str(o) for o in opts]
    return out


REF_RE = re.compile(r'"[^"]*"|(\$?)([A-Z]{1,3})(\$?)(\d+)')


def _rgb(c):
    # ponytail: only explicit ARGB colors; theme/indexed colors are skipped
    return f"#{c.rgb[-6:]}" if c is not None and isinstance(c.rgb, str) and c.rgb[-6:] != "000000" else None


def cond_styles(ws, max_row):
    """Map (row, col) -> css for expression-type conditional formats.

    Supports cell refs, string/number literals, = <> < > <= >=, AND/OR/NOT.
    Per property, the rule with the lowest priority number wins (Excel's order).
    """
    picks = {}
    for cf in ws.conditional_formatting:
        for rule in cf.rules:
            if rule.type != "expression" or not rule.formula or not rule.dxf:
                continue
            d = rule.dxf
            props = {}
            if d.fill is not None:
                props["background"] = _rgb(d.fill.bgColor) or _rgb(d.fill.fgColor)
            if d.font is not None:
                props["color"] = _rgb(d.font.color)
                props["font-weight"] = "600" if d.font.b else None
            props = {k: v for k, v in props.items() if v}
            if not props:
                continue
            for rng in str(cf.sqref).split():
                c1, r1, c2, r2 = range_boundaries(rng)
                for r in range(r1, min(r2, max_row) + 1):
                    for c in range(c1, c2 + 1):
                        if _eval(ws, rule.formula[0], r - r1, c - c1):
                            for k, v in props.items():
                                cur = picks.setdefault((r, c), {}).get(k)
                                if cur is None or rule.priority < cur[0]:
                                    picks[(r, c)][k] = (rule.priority, v)
    return {rc: ";".join(f"{k}:{v}" for k, (_, v) in p.items()) for rc, p in picks.items()}


def _eval(ws, formula, dr, dc):
    def ref(m):
        if m.group(2) is None:
            return m.group(0)  # string literal
        col = column_index_from_string(m.group(2)) + (0 if m.group(1) else dc)
        row = int(m.group(4)) + (0 if m.group(3) else dr)
        v = ws.cell(row, col).value
        return repr("" if v is None else v)

    expr = REF_RE.sub(ref, formula)
    expr = expr.replace("<>", "!=")
    expr = re.sub(r'(?<![<>!=])=(?!=)', "==", expr)
    expr = re.sub(r"\b(AND|OR|NOT)\(", lambda m: f"_{m.group(1).lower()}(", expr)
    try:
        return bool(eval(expr, {"__builtins__": {}}, {
            "_and": lambda *a: all(a), "_or": lambda *a: any(a), "_not": lambda a: not a}))
    except Exception:
        return False  # unsupported formula: no color rather than a crash


def fmt(v):
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    return "" if v is None else str(v)


def render():
    wb = openpyxl.load_workbook(PATH)
    tabs, bodies = [], []
    for i, ws in enumerate(wb):
        dd = dropdowns(wb, ws)
        # ponytail: trailing empty rows trimmed, plus 5 blank rows for adding
        last = max((c.row for row in ws.iter_rows() for c in row if c.value is not None), default=1)
        cols = ws.max_column
        cs = cond_styles(ws, last + 5)
        rows = []
        for r in range(1, last + 6):
            tds = [f"<th>{r}</th>"]
            for c in range(1, cols + 1):
                v = html.escape(fmt(ws.cell(r, c).value))
                attr = f'data-s="{i}" data-r="{r}" data-c="{c}"'
                st = f' style="{cs[(r, c)]}"' if (r, c) in cs else ""
                if (r, c) in dd:
                    opts = "".join(
                        f'<option{" selected" if o == fmt(ws.cell(r, c).value) else ""}>{html.escape(o)}</option>'
                        for o in [""] + dd[(r, c)]
                    )
                    tds.append(f"<td{st}><select {attr}>{opts}</select></td>")
                else:
                    tds.append(f"<td contenteditable{st} {attr}>{v}</td>")
            rows.append(f"<tr{' class=hd' if r == 1 else ''}>{''.join(tds)}</tr>")
        tabs.append(f'<button onclick="show({i})" id="t{i}">{html.escape(ws.title)}</button>')
        bodies.append(f'<table id="s{i}" hidden>{"".join(rows)}</table>')
    return PAGE.replace("{{TITLE}}", html.escape(PATH.split("/")[-1])).replace(
        "{{TABS}}", "".join(tabs)).replace("{{BODIES}}", "".join(bodies))


def save(changes):
    wb = openpyxl.load_workbook(PATH)
    sheets = wb.worksheets
    for ch in changes:
        v = ch["v"].strip()
        if v == "":
            v = None
        elif DATE_RE.match(v):
            v = datetime.date.fromisoformat(v)
        elif re.fullmatch(r"-?\d+(\.\d+)?", v) and not v.startswith("0"):
            v = float(v) if "." in v else int(v)
        cell = sheets[ch["s"]].cell(ch["r"], ch["c"])
        cell.value = v
        if isinstance(v, datetime.date):
            cell.number_format = "yyyy-mm-dd"
    wb.save(PATH)


class H(BaseHTTPRequestHandler):
    def do_HEAD(self):
        self.send_response(200)
        self.end_headers()

    def do_GET(self):
        try:
            body = render().encode()
        except Exception as e:  # show why instead of a blank page
            body = f"<pre>열기 실패: {html.escape(str(e))}</pre>".encode()
        self.send_response(200)
        self.send_header("Content-Type", "text/html; charset=utf-8")
        self.end_headers()
        self.wfile.write(body)

    def do_POST(self):
        try:
            save(json.loads(self.rfile.read(int(self.headers["Content-Length"]))))
            self.send_response(204)
        except Exception as e:  # surface save errors (e.g. file locked) to the page
            self.send_response(500)
            self.end_headers()
            self.wfile.write(str(e).encode())
            return
        self.end_headers()


PAGE = """<!doctype html><html lang="ko"><meta charset="utf-8"><title>{{TITLE}}</title>
<style>
body{font:13px -apple-system,sans-serif;margin:0;background:#fff;color:#222}
header{position:sticky;top:0;background:#f4f4f5;padding:8px;display:flex;gap:6px;align-items:center;border-bottom:1px solid #ddd;z-index:2}
header button{padding:4px 10px;border:1px solid #ccc;background:#fff;border-radius:4px;cursor:pointer}
header button.on{background:#222;color:#fff}
#save{margin-left:auto;background:#2563eb;color:#fff;border:0}
#msg{color:#666}
table{border-collapse:collapse;margin:8px}
td,th{border:1px solid #ddd;padding:3px 6px;min-width:60px;max-width:420px;vertical-align:top;white-space:pre-wrap}
th{background:#f4f4f5;color:#888;font-weight:normal}
tr.hd td{background:#eef2ff;font-weight:600}
td.dirty,select.dirty{background:#fef9c3}
select{border:0;background:transparent;font:inherit;color:inherit;width:100%}
</style>
<header>{{TABS}}<span id="msg"></span><button id="save" onclick="save()">저장 (⌘S)</button></header>
{{BODIES}}
<script>
const dirty=new Map();
function show(i){document.querySelectorAll('table').forEach((t,j)=>t.hidden=i!==j);
document.querySelectorAll('header button[id^=t]').forEach((b,j)=>b.classList.toggle('on',i===j))}
function mark(e){const t=e.target;if(!t.dataset.r)return;t.classList.add('dirty');
dirty.set(`${t.dataset.s},${t.dataset.r},${t.dataset.c}`,t);msg.textContent=`변경 ${dirty.size}칸`}
document.addEventListener('input',mark);document.addEventListener('change',mark);
async function save(){if(!dirty.size)return;
const body=[...dirty.values()].map(t=>({s:+t.dataset.s,r:+t.dataset.r,c:+t.dataset.c,v:t.tagName==='SELECT'?t.value:t.innerText}));
const r=await fetch('/',{method:'POST',body:JSON.stringify(body)});
if(r.ok){dirty.forEach(t=>t.classList.remove('dirty'));dirty.clear();location.reload()}
else msg.textContent='저장 실패: '+await r.text()}
document.addEventListener('keydown',e=>{if((e.metaKey||e.ctrlKey)&&e.key==='s'){e.preventDefault();save()}});
window.onbeforeunload=()=>dirty.size?true:undefined;
show(0);
</script></html>"""

if __name__ == "__main__":
    print(f"http://127.0.0.1:{PORT}  ({PATH})")
    ThreadingHTTPServer(("127.0.0.1", PORT), H).serve_forever()

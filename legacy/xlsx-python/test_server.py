"""Run: python3 test_server.py"""
import sys

sys.argv = ["x", "unused.xlsx"]
import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill

import server

wb = openpyxl.Workbook()
ws = wb.active
ws.append(["owner", "manager", "grade"])
ws.append(["a", "a", "1급"])
ws.append(["a", "b", "1급"])
ws.append(["", "", "1급"])
red = PatternFill(bgColor="F6DEDC")
ws.conditional_formatting.add("B2:B10", FormulaRule(formula=['AND($C2="1급",$A2<>"",$A2=$B2)'], fill=red, font=Font(color="8C1D18")))

cs = server.cond_styles(ws, 10)
assert cs[(2, 2)] == "background:#F6DEDC;color:#8C1D18", cs
assert (3, 2) not in cs and (4, 2) not in cs, cs
assert server._eval(ws, 'NOT($C2="2급")', 0, 0)
assert not server._eval(ws, "BOGUS(", 0, 0)
print("ok")
